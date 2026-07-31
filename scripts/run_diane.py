"""Pilote one-shot : adapte les CV de Diane (PDF + .docx) sur l'orientation cible."""
from pathlib import Path

import docx

from src.analysis.parsing import parse_pdf, parse_text
from src.analysis.pipeline.orchestrate import generate
from src.config import OUTPUT_DIR
from src.rendering.constraints import default_constraints
from src.rendering.pptx_renderer import render

INPUT = Path("input/Diane")

PDF_CVS = [
    "CV Diane Ademe pilotage.pdf",
    "CV Diane Maurin revit audit.pdf",
    "CV Diane versions Agile & IA.pdf",
    "Profile.pdf",
]
DOCX_CV = "Compétences transverses Diane Maurin.docx"

AO_TEXT = (
    "Mission orientée secteur énergie au sein d'une DSI. Pilotage de "
    "transformations complexes : appui au pilotage de programmes et projets, "
    "conduite du changement, déploiement d'outils et méthodes de transformation. "
    "Mobilisation et animation de collectifs (ateliers, communautés, équipes). "
    "Conception de supports pédagogiques et communicants. Profil recherché : "
    "capacité à structurer, à embarquer les parties prenantes et à outiller la "
    "transformation dans un environnement SI du secteur de l'énergie."
)


def docx_text(path: Path) -> str:
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def main() -> None:
    cvs = [parse_pdf(INPUT / name) for name in PDF_CVS]
    cvs.append(parse_text(docx_text(INPUT / DOCX_CV), DOCX_CV))

    ao = parse_text(AO_TEXT, "orientation-diane.txt")
    constraints = default_constraints()

    print(f"[run] {len(cvs)} sources CV, AO={ao.source_name}")
    adapted = generate(cvs, ao, constraints)
    print(f"[run] couverture besoins : {adapted.coverage.score_global}/100")
    print(f"      domaines retenus    : {len(adapted.experiences)}")
    print(f"      compétences retenues: {len(adapted.skills)}")
    if adapted.alertes_completude:
        print(f"      alertes complétude  : {len(adapted.alertes_completude)}")
        for a in adapted.alertes_completude:
            print(f"        - {a}")

    out = render(adapted, constraints)
    if out.truncations:
        print(f"      troncatures ({len(out.truncations)}) :")
        for t in out.truncations:
            print(f"        - {t}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    pptx_path = OUTPUT_DIR / "CV-Diane-energie-DSI.pptx"
    json_path = OUTPUT_DIR / "CV-Diane-energie-DSI.json"
    pptx_path.write_bytes(out.pptx_bytes)
    json_path.write_text(adapted.model_dump_json(indent=2), encoding="utf-8")
    print(f"[run] PPTX écrit : {pptx_path} ({len(out.pptx_bytes)} octets)")
    print(f"[run] JSON écrit : {json_path}")


if __name__ == "__main__":
    main()
